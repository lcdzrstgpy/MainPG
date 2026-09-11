# -*- coding: utf-8 -*-
"""妙手 Temu 导入模板导出：模板缓存、列映射与逐行取值。

产品处理与 POD 定制都从这里生成妙手表格，保证两个入口的字段口径完全一致：

- 产品处理：``domain/workbooks.py`` 先把商品结果转成店小秘导出行，再调用本模块；
- POD 定制：``pod_customization/export.py`` 已有店小秘行，直接调用本模块。

官方模板与模板缓存仍存放在 ``modules/product_processing/templates``（打包清单与
.gitignore 均指向该目录），本模块只负责读取。
"""

from __future__ import annotations

import os
import threading
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

# 无发货时效数据时的默认值（店小秘/妙手共用）。
DEFAULT_SHIP_DAYS = 2

MS_KIND_APPAREL = "apparel"
MS_KIND_GENERAL = "general"
MS_KIND_OPTIONS = (MS_KIND_APPAREL, MS_KIND_GENERAL)

# 模板目录保持原地：打包清单（workbench.spec）与 .gitignore 均绑定该路径。
MS_TEMPLATE_DIR = (
    Path(__file__).resolve().parents[1] / "modules" / "product_processing" / "templates"
)
MS_TEMPLATE_FILES = {
    MS_KIND_APPAREL: "妙手Temu导入模板-服饰类模板.xlsx",
    MS_KIND_GENERAL: "妙手Temu导入模板-非服饰类模板.xlsx",
}

MS_GENERATED_DIR_NAME = "generated"
MS_SLIM_TEMPLATE_FILES = {
    MS_KIND_APPAREL: "妙手-服饰类.sheet1.xlsx",
    MS_KIND_GENERAL: "妙手-非服饰类.sheet1.xlsx",
}
MS_REFERENCE_FILE = "妙手参考表-类目ID与包装清单.xlsx"
MS_MAIN_SHEET = "Sheet1"

_MS_TEMPLATE_LOCK = threading.Lock()
_MS_REFERENCE_LOCK = threading.Lock()

# 系统无数据/固定策略列的默认值。
MS_ORIGIN_DEFAULT = "中国-浙江省"
MS_CUSTOMIZED_DEFAULT = "否"
MS_SENSITIVE_DEFAULT = "否"

# 服饰类模板（Sheet1，1-based 列号），列结构见模板表头：类目ID/主编号/标题/英文标题/...
_MS_APPAREL_COLUMNS = {
    "category_id": 2,
    "main_no": 3,
    "title": 4,
    "title_en": 5,
    "description": 6,
    "ship_days": 7,
    "origin": 8,
    "made_in": 9,
    "external_url": 10,
    "material_image": 11,
    "customized": 12,
    "spec_name_1": 13,
    "spec_value_1": 14,
    "spec_name_2": 15,
    "spec_value_2": 16,
    "color_images": 17,
    "main_sku_no": 18,
    "declared_price": 19,
    "suggested_price": 20,
    "length_cm": 21,
    "width_cm": 22,
    "height_cm": 23,
    "weight_g": 24,
    "stock": 25,
    "platform_sku": 26,
    "sensitive": 27,
    "sensitive_value": 28,
    "code_type": 33,
    "code": 34,
    "sku_class_type": 35,
    "sku_class_count": 36,
    "sku_class_unit": 37,
    "independent_pack": 38,
    "pack_list": 39,
    "pack_list_count": 40,
    "video": 41,
    "manual": 42,
    "supply_url": 43,
}

# 非服饰类模板（Sheet1，1-based 列号）：列位与服饰类整体平移，图片列结构不同。
_MS_GENERAL_COLUMNS = {
    "category_id": 2,
    "main_no": 3,
    "title": 4,
    "title_en": 5,
    "description": 6,
    "ship_days": 7,
    "main_sku_no": 8,
    "origin": 9,
    "made_in": 10,
    "external_url": 11,
    "carousel_images": 12,
    "material_image": 13,
    "customized": 14,
    "spec_name_1": 15,
    "spec_value_1": 16,
    "spec_name_2": 17,
    "spec_value_2": 18,
    "preview_image": 19,
    "declared_price": 20,
    "suggested_price": 21,
    "length_cm": 22,
    "width_cm": 23,
    "height_cm": 24,
    "weight_g": 25,
    "stock": 26,
    "platform_sku": 27,
    "sensitive": 28,
    "sensitive_value": 29,
    "code_type": 34,
    "code": 35,
    "sku_class_type": 36,
    "sku_class_count": 37,
    "sku_class_unit": 38,
    "independent_pack": 39,
    "pack_list": 40,
    "pack_list_count": 41,
    "video": 42,
    "manual": 43,
    "supply_url": 44,
}


def miaoshou_columns(kind: str) -> Mapping[str, int]:
    """按模板类型返回列号映射（1-based）。"""
    if kind == MS_KIND_APPAREL:
        return _MS_APPAREL_COLUMNS
    if kind == MS_KIND_GENERAL:
        return _MS_GENERAL_COLUMNS
    raise ValueError(f"unsupported miaoshou template kind: {kind}")


def _miaoshou_generated_dir(template_dir: Path) -> Path:
    return template_dir / MS_GENERATED_DIR_NAME


def ensure_miaoshou_reference_workbook(*, template_dir: Path | None = None) -> Path:
    """生成/更新妙手参考表文件（类目ID + 包装清单），供人工查阅。

    这两张表原本在官方模板里，导出时一并携带会让 openpyxl 每次多花约 8 秒。
    模板更新或参考表缺失时重建一次；耗时较长（约 8 秒），导出路径通过后台线程调用。
    """
    base = template_dir or MS_TEMPLATE_DIR
    full = base / MS_TEMPLATE_FILES[MS_KIND_GENERAL]
    if not full.is_file():
        raise ValueError(f"miaoshou template not bundled: {full}")
    reference = base / MS_REFERENCE_FILE
    with _MS_REFERENCE_LOCK:
        try:
            if reference.is_file() and reference.stat().st_mtime >= full.stat().st_mtime:
                return reference
        except OSError:
            pass
        workbook = load_workbook(full)
        try:
            for name in list(workbook.sheetnames):
                if name == MS_MAIN_SHEET:
                    workbook.remove(workbook[name])
            reference.parent.mkdir(parents=True, exist_ok=True)
            temporary = reference.with_name(f".{reference.name}.tmp")
            workbook.save(temporary)
        finally:
            workbook.close()
        os.replace(temporary, reference)
        return reference


def ensure_miaoshou_reference_async(*, template_dir: Path | None = None) -> None:
    """后台生成参考表，避免拖慢当前导出请求（失败静默，下次导出会重试）。"""

    def _run() -> None:
        try:
            ensure_miaoshou_reference_workbook(template_dir=template_dir)
        except Exception:  # noqa: BLE001 - 参考表生成失败不应影响导出
            return

    threading.Thread(target=_run, name="miaoshou-reference-builder", daemon=True).start()


def miaoshou_slim_template_path(kind: str, *, template_dir: Path | None = None) -> Path:
    """精简模板（只含 Sheet1）路径；缺失或官方模板更新后自动重建。

    官方模板的两张参考表（类目ID/包装清单）只供人工查阅，但会让 openpyxl 每次
    加载/保存多花约 8 秒。这里生成只含 Sheet1 的缓存副本，导出时读它即可；
    参考表则抽成独立文件（ensure_miaoshou_reference_workbook）另行查阅。
    """
    if kind not in MS_TEMPLATE_FILES:
        raise ValueError(f"unsupported miaoshou template kind: {kind}")
    base = template_dir or MS_TEMPLATE_DIR
    full = base / MS_TEMPLATE_FILES[kind]
    if not full.is_file():
        raise ValueError(f"miaoshou template not bundled: {full}")
    slim = _miaoshou_generated_dir(base) / MS_SLIM_TEMPLATE_FILES[kind]
    with _MS_TEMPLATE_LOCK:
        try:
            if slim.is_file() and slim.stat().st_mtime >= full.stat().st_mtime:
                return slim
        except OSError:
            pass
        workbook = load_workbook(full)
        try:
            for name in list(workbook.sheetnames):
                if name != MS_MAIN_SHEET:
                    workbook.remove(workbook[name])
            slim.parent.mkdir(parents=True, exist_ok=True)
            temporary = slim.with_name(f".{slim.name}.tmp")
            workbook.save(temporary)
        except OSError:
            # 安装目录不可写时退化为直接读官方模板：慢一些但导出仍可用。
            return full
        finally:
            workbook.close()
        try:
            os.replace(temporary, slim)
        except OSError:
            return full
    # 参考表与精简模板同源同代，重建精简模板时顺带后台补齐参考表。
    ensure_miaoshou_reference_async(template_dir=template_dir)
    return slim


def miaoshou_row_values(
    dxm_row: Sequence[Any],
    kind: str,
    *,
    stock: Any,
    description: str = "",
    main_no: str | None = None,
    weight_g: Any = None,
    ship_days: Any = None,
    target_language: str = "",
) -> dict[int, Any]:
    """把一行店小秘导出结果转成妙手模板的列值（列号 -> 值）。

    调用方负责准备两个「两个入口口径不同」的值：

    - ``description``：妙手描述列不支持 HTML，需由调用方拼成「正文 + 详情图 URL 逐行」；
    - ``stock``：妙手要求库存为 >=0 的整数（0 也必须显式写出，留空会被判导入失败）。

    其余字段直接取店小秘行，保证与店小秘导出同源：

    - 主编号：默认取店小秘「产品货号」列，产品处理侧改传商品级 SKC/货号；
    - 重量：默认取店小秘重量列，产品处理侧改传「不做取整与封顶」的真实重量；
    - 规格名称2/值2 为妙手必填：无第二规格轴时兜底「规格 / Standard」（西语站 Estándar）。
    """
    columns = miaoshou_columns(kind)

    values: dict[int, Any] = {}
    title = str(dxm_row[0] or "").strip()
    values[columns["title"]] = title
    values[columns["title_en"]] = title
    values[columns["main_no"]] = (
        str(main_no).strip() if main_no is not None else str(dxm_row[3] or "").strip()
    )
    values[columns["description"]] = description

    # 规格名/值（dxm 行 4..7 列）。规格名称2/值2 在妙手模板中为必填：有第二规格轴时
    # 用真实值；没有时兜底为「规格 / Standard」，避免模板因必填列为空被拒。
    values[columns["spec_name_1"]] = dxm_row[4]
    values[columns["spec_value_1"]] = dxm_row[5]
    spec_name_2 = dxm_row[6] if dxm_row[6] else ""
    spec_value_2 = dxm_row[7] if dxm_row[7] else ""
    if not (spec_name_2 and spec_value_2):
        spec_name_2 = "规格"
        spec_value_2 = (
            "Estándar" if str(target_language or "").strip().casefold() == "es" else "Standard"
        )
    values[columns["spec_name_2"]] = spec_name_2
    values[columns["spec_value_2"]] = spec_value_2

    # 图片：与店小秘导出同一批最终图片（均已是公网 https）。
    carousel_text = str(dxm_row[18] or "").strip()
    main_image = str(dxm_row[8] or "").strip()
    if not main_image and carousel_text:
        main_image = carousel_text.splitlines()[0]
    if kind == MS_KIND_APPAREL:
        values[columns["color_images"]] = carousel_text
        values[columns["material_image"]] = main_image
    else:
        values[columns["carousel_images"]] = carousel_text
        values[columns["material_image"]] = main_image
        values[columns["preview_image"]] = main_image

    # 价格/物流/库存。
    values[columns["declared_price"]] = dxm_row[9]
    values[columns["suggested_price"]] = dxm_row[23] if dxm_row[23] not in (None, "") else ""
    values[columns["length_cm"]] = dxm_row[11]
    values[columns["width_cm"]] = dxm_row[12]
    values[columns["height_cm"]] = dxm_row[13]
    values[columns["weight_g"]] = dxm_row[14] if weight_g is None else weight_g
    values[columns["stock"]] = stock
    if ship_days is None:
        ship_days = dxm_row[25] if dxm_row[25] not in (None, "") else DEFAULT_SHIP_DAYS
    values[columns["ship_days"]] = ship_days

    # 系统无数据/固定策略列。
    values[columns["origin"]] = MS_ORIGIN_DEFAULT
    values[columns["customized"]] = MS_CUSTOMIZED_DEFAULT
    values[columns["sensitive"]] = MS_SENSITIVE_DEFAULT
    values[columns["external_url"]] = str(dxm_row[17] or "").strip()
    # 货源链接（选填）店小秘无对应数据，妙手支持为空，这里留空由用户自行补充。
    values[columns["supply_url"]] = ""
    return values


def write_miaoshou_workbook(
    rows: Sequence[Mapping[int, Any]],
    kind: str,
    destination: Path | BytesIO,
    *,
    template_dir: Path | None = None,
) -> int:
    """按妙手导入模板写出工作簿：保留模板 Sheet1 结构，逐行写入已展开的列值。

    ``rows`` 是 ``miaoshou_row_values`` 的结果序列（每个元素 = 一个 SKU 行）。
    返回实际写入的数据行数。模板第 1 行表头、第 2 行字段说明保留，第 3 行起的
    示例数据将被清空后重写。
    """
    if kind not in MS_TEMPLATE_FILES:
        raise ValueError(f"unsupported miaoshou template kind: {kind}")
    slim_path = miaoshou_slim_template_path(kind, template_dir=template_dir)

    workbook = load_workbook(slim_path)
    try:
        sheet = workbook[MS_MAIN_SHEET]
        # 移除模板示例行上残留的合并单元格后，删除第 3 行起全部示例内容。
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        if sheet.max_row and sheet.max_row > 2:
            sheet.delete_rows(3, sheet.max_row - 2)

        row_number = 3
        for values in rows:
            for column, value in values.items():
                sheet.cell(row=row_number, column=column, value=value)
            row_number += 1

        target = Path(destination) if isinstance(destination, (str, os.PathLike)) else destination
        if isinstance(target, Path):
            target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
    finally:
        workbook.close()
    return row_number - 3


def build_miaoshou_workbook_bytes(
    rows: Sequence[Mapping[int, Any]], kind: str, *, template_dir: Path | None = None
) -> bytes:
    """生成妙手工作簿字节内容（不落盘），供 HTTP 直接下载。"""
    buffer = BytesIO()
    write_miaoshou_workbook(rows, kind, buffer, template_dir=template_dir)
    return buffer.getvalue()
