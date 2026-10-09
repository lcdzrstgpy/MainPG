from __future__ import annotations

import io
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from wh_local.modules.pod_customization.contracts import (
    ListingFields,
    ReplicaBatchCreate,
    ReplicaTargetCreate,
)
from wh_local.modules.pod_customization.replica_context import style_product_context
from wh_local.modules.pod_customization.repository import PodRepositoryError
from wh_local.modules.pod_customization.service import PodCustomizationService
from wh_local.session import Actor


class NeverCalledRuntime:
    def submit(self, *_args, **_kwargs):
        raise AssertionError("export must not call AI")


class NeverCalledTitleRuntime:
    configured = True

    def submit(self, *_args, **_kwargs):
        raise AssertionError("export must not call title AI")


_IMAGE_ROLES = ("hero", "detail_a", "detail_b", "lifestyle")


def _png() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (320, 240), "#eee9df").save(output, "PNG")
    return output.getvalue()


def _actor(user_id: str = "operator-1") -> Actor:
    return Actor(id=user_id, username=user_id, role="operator", workspace_id="workspace-a")


def _service(tmp_path: Path) -> PodCustomizationService:
    return PodCustomizationService(
        tmp_path / "workbench.sqlite3",
        tmp_path / "assets",
        NeverCalledRuntime(),
        title_runtime=NeverCalledTitleRuntime(),
        start_workers=False,
    )


def _listing_fields(
    *,
    category: str,
    suggested_price: float,
    skus: list[dict[str, object]],
    display_unit: str = "cm",
) -> ListingFields:
    return ListingFields(
        suggested_price_usd=suggested_price,
        category_name=category,
        skus=[
            {
                "name": sku["name"],
                "declared_price": sku["declared_price"],
                "weight_g": sku["weight_g"],
            }
            for sku in skus
        ],
        spec_card={
            "display_unit": display_unit,
            "cells": [
                ["SKU", "Length", "Width", "Height"],
                *[
                    [
                        str(sku["name"]),
                        str(sku["length_cm"]),
                        str(sku["width_cm"]),
                        str(sku["height_cm"]),
                    ]
                    for sku in skus
                ],
            ],
        },
    )


def _upload(service: PodCustomizationService, actor: Actor, role: str, filename: str) -> str:
    return service.upload_replica_image(actor, role=role, filename=filename, content=_png())["asset_id"]


def _target(asset_id: str, product_name: str, listing: ListingFields) -> ReplicaTargetCreate:
    return ReplicaTargetCreate(
        target_asset_id=asset_id,
        product_name=product_name,
        listing_fields=listing,
    )


def _request(source_id: str, targets: list[ReplicaTargetCreate], client_request_id: str) -> ReplicaBatchCreate:
    return ReplicaBatchCreate(
        client_request_id=client_request_id,
        source_asset_id=source_id,
        targets=targets,
    )


def _complete_style(
    service: PodCustomizationService,
    batch_id: str,
    style_index: int,
    prefix: str,
    *,
    urls: tuple[str, str, str, str] | None = None,
) -> tuple[str, str, str, str]:
    urls = urls or tuple(
        f"https://images.example.com/replica/{prefix}/{role}.png" for role in _IMAGE_ROLES
    )
    with sqlite3.connect(service.database_path) as connection:
        rows = connection.execute(
            """SELECT result_id, variant_index FROM pod_customization_style_grid_results
               WHERE batch_id = ? AND style_index = ? ORDER BY variant_index""",
            (batch_id, style_index),
        ).fetchall()
        for (result_id, _variant_index), role, url in zip(rows, _IMAGE_ROLES, urls, strict=True):
            connection.execute(
                """UPDATE pod_customization_style_grid_results
                   SET status = 'completed', pattern_asset_id = 'pattern', composite_asset_id = 'composite'
                   WHERE result_id = ?""",
                (result_id,),
            )
            connection.execute(
                """INSERT INTO pod_customization_style_grid_publications
                   (result_id, role, public_url, updated_at) VALUES (?, ?, ?, 'now')""",
                (result_id, role, url),
            )
    return urls


def _settle(service: PodCustomizationService, batch_id: str, status: str = "completed") -> None:
    with sqlite3.connect(service.database_path) as connection:
        connection.execute(
            "UPDATE pod_customization_batches SET status = ? WHERE batch_id = ?", (status, batch_id)
        )


def _complete_listing_title(service: PodCustomizationService, batch_id: str, style_index: int) -> None:
    with sqlite3.connect(service.database_path) as connection:
        connection.execute(
            "UPDATE pod_customization_style_titles SET status = 'completed' WHERE batch_id = ? AND style_index = ?",
            (batch_id, style_index),
        )


def _two_product_batch(
    service: PodCustomizationService,
    actor: Actor,
    *,
    listing_a: ListingFields,
    listing_b: ListingFields,
    client_request_id: str,
) -> dict:
    source_id = _upload(service, actor, "source", "source.png")
    target_a = _upload(service, actor, "target", "cushion.png")
    target_b = _upload(service, actor, "target", "basket.png")
    return service.create_replica_batch(
        actor,
        _request(
            source_id,
            [
                _target(target_a, "抱枕", listing_a),
                _target(target_b, "收纳篮", listing_b),
            ],
            client_request_id=client_request_id,
        ),
        enqueue=False,
    )


def test_replica_export_expands_each_product_sku_without_cross_style_fallback(tmp_path: Path) -> None:
    service = _service(tmp_path)
    actor = _actor()
    listing_a = _listing_fields(
        category="家居收纳 > 抱枕",
        suggested_price=19.99,
        skus=[
            {"name": "A-40", "declared_price": 6.0, "weight_g": 300.0, "length_cm": 40, "width_cm": 40, "height_cm": 10},
            {"name": "A-50", "declared_price": 8.0, "weight_g": 500.0, "length_cm": 50, "width_cm": 50, "height_cm": 15},
        ],
        display_unit="cm",
    )
    listing_b = _listing_fields(
        category="家居收纳 > 收纳篮",
        suggested_price=12.5,
        skus=[
            {"name": "B-30", "declared_price": 4.0, "weight_g": 200.0, "length_cm": 30, "width_cm": 30, "height_cm": 20},
        ],
        display_unit="in",
    )
    batch = _two_product_batch(
        service, actor, listing_a=listing_a, listing_b=listing_b, client_request_id="req-export-two"
    )
    batch_id = batch["id"]

    _complete_style(service, batch_id, 1, "a")
    _complete_style(service, batch_id, 2, "b")
    _settle(service, batch_id)
    _complete_listing_title(service, batch_id, 1)
    _complete_listing_title(service, batch_id, 2)
    service.repository.upsert_style_copy(
        batch_id, actor.workspace_id, actor.id, 1,
        title="Cushion Cover Floral", english_title="Cushion Cover", description="A cushion cover.",
    )
    service.repository.upsert_style_copy(
        batch_id, actor.workspace_id, actor.id, 2,
        title="Storage Basket Striped", english_title="Storage Basket", description="A storage basket.",
    )

    stored = service.repository.get_batch(batch_id, actor.workspace_id, actor.id)
    # 按款读取字段：类目不得回退首款。
    assert style_product_context(stored, 1)["listing_fields"]["category_name"] == "家居收纳 > 抱枕"
    assert style_product_context(stored, 2)["listing_fields"]["category_name"] == "家居收纳 > 收纳篮"

    exported = service.export_dianxiaomi(actor, batch_id)
    assert exported.exported_style_count == 2
    assert exported.skipped_style_count == 0

    workbook = load_workbook(io.BytesIO(exported.content), data_only=True)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()

    assert len(rows) == 3
    # 商品码按款号、SKU 码按款内名：款间不合并、不串行。
    assert [row[3] for row in rows] == ["POD-001", "POD-001", "POD-002"]
    assert [row[5] for row in rows] == ["A-40", "A-50", "B-30"]
    assert [row[10] for row in rows] == ["A-40", "A-50", "B-30"]
    # 标题归对应产品。
    assert [row[0] for row in rows] == [
        "Cushion Cover Floral", "Cushion Cover Floral", "Storage Basket Striped",
    ]
    # 重量归对应产品。
    assert [row[14] for row in rows] == [300, 500, 200]
    # 建议售价归对应产品。
    assert [row[23] for row in rows] == [19.99, 19.99, 12.5]
    # 尺寸：款 1 保留 cm，款 2 从 cm 换算成 in（30/2.54≈11.81，20/2.54≈7.87）。
    assert [row[11:14] for row in rows] == [
        (40.0, 40.0, 10.0),
        (50.0, 50.0, 15.0),
        (11.81, 11.81, 7.87),
    ]
    # 图片角色归对应产品。
    assert [row[8] for row in rows] == [
        "https://images.example.com/replica/a/lifestyle.png",
        "https://images.example.com/replica/a/lifestyle.png",
        "https://images.example.com/replica/b/lifestyle.png",
    ]
    assert [row[19] for row in rows] == [
        "https://images.example.com/replica/a/hero.png",
        "https://images.example.com/replica/a/hero.png",
        "https://images.example.com/replica/b/hero.png",
    ]


def test_replica_export_partial_failure_exports_ready_product_only(tmp_path: Path) -> None:
    service = _service(tmp_path)
    actor = _actor()
    listing = _listing_fields(
        category="家居收纳 > 抱枕",
        suggested_price=19.99,
        skus=[
            {"name": "A-40", "declared_price": 6.0, "weight_g": 300.0, "length_cm": 40, "width_cm": 40, "height_cm": 10},
        ],
    )
    batch = _two_product_batch(
        service, actor, listing_a=listing, listing_b=listing, client_request_id="req-partial"
    )
    batch_id = batch["id"]

    _complete_style(service, batch_id, 1, "a")
    # 款 2 首图为非公网 http，整款四图不合格 → 应跳过。
    _complete_style(
        service, batch_id, 2, "b",
        urls=(
            "http://images.example.com/replica/b/hero.png",
            "https://images.example.com/replica/b/detail_a.png",
            "https://images.example.com/replica/b/detail_b.png",
            "https://images.example.com/replica/b/lifestyle.png",
        ),
    )
    _settle(service, batch_id, "partial_failure")
    _complete_listing_title(service, batch_id, 1)
    _complete_listing_title(service, batch_id, 2)
    service.repository.upsert_style_copy(
        batch_id, actor.workspace_id, actor.id, 1,
        title="Cushion Cover Floral", english_title="Cushion Cover", description="A cushion cover.",
    )
    service.repository.upsert_style_copy(
        batch_id, actor.workspace_id, actor.id, 2,
        title="Storage Basket Striped", english_title="Storage Basket", description="A storage basket.",
    )

    exported = service.export_dianxiaomi(actor, batch_id)
    assert exported.exported_style_count == 1
    assert exported.skipped_style_count == 1
    workbook = load_workbook(io.BytesIO(exported.content), data_only=True)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    assert [row[3] for row in rows] == ["POD-001"]
    assert rows[0][0] == "Cushion Cover Floral"


def test_replica_export_excludes_unselected_product(tmp_path: Path) -> None:
    service = _service(tmp_path)
    actor = _actor()
    listing = _listing_fields(
        category="家居收纳 > 抱枕",
        suggested_price=19.99,
        skus=[
            {"name": "A-40", "declared_price": 6.0, "weight_g": 300.0, "length_cm": 40, "width_cm": 40, "height_cm": 10},
        ],
    )
    batch = _two_product_batch(
        service, actor, listing_a=listing, listing_b=listing, client_request_id="req-unselected"
    )
    batch_id = batch["id"]

    _complete_style(service, batch_id, 1, "a")
    _complete_style(service, batch_id, 2, "b")
    _settle(service, batch_id)
    _complete_listing_title(service, batch_id, 1)
    _complete_listing_title(service, batch_id, 2)
    for style_index, title in ((1, "Cushion Cover Floral"), (2, "Storage Basket Striped")):
        service.repository.upsert_style_copy(
            batch_id, actor.workspace_id, actor.id, style_index,
            title=title, english_title=title, description=f"{title} description.",
        )

    service.set_style_export_selection(actor, batch_id, 2, selected=False)
    exported = service.export_dianxiaomi(actor, batch_id)
    assert exported.exported_style_count == 1
    assert exported.skipped_style_count == 1
    workbook = load_workbook(io.BytesIO(exported.content), data_only=True)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    assert [row[3] for row in rows] == ["POD-001"]
    assert rows[0][0] == "Cushion Cover Floral"


def test_replica_export_no_exportable_product_blocks(tmp_path: Path) -> None:
    service = _service(tmp_path)
    actor = _actor()
    listing = _listing_fields(
        category="家居收纳 > 抱枕",
        suggested_price=19.99,
        skus=[
            {"name": "A-40", "declared_price": 6.0, "weight_g": 300.0, "length_cm": 40, "width_cm": 40, "height_cm": 10},
        ],
    )
    batch = _two_product_batch(
        service, actor, listing_a=listing, listing_b=listing, client_request_id="req-empty"
    )
    batch_id = batch["id"]
    # 不完成任何四图，仅落一个文案。
    _settle(service, batch_id)
    service.repository.upsert_style_copy(
        batch_id, actor.workspace_id, actor.id, 1,
        title="Title", english_title="English", description="Description",
    )

    payload = service.get_batch(actor, batch_id)
    assert payload["dianxiaomi_export"]["block_reason"] == "no_exportable_styles"
    assert payload["dianxiaomi_export"]["ready"] is False
    with pytest.raises(PodRepositoryError, match="可导出") as raised:
        service.export_dianxiaomi(actor, batch_id)
    assert raised.value.status_code == 409